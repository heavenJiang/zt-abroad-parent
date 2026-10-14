"""M2-07 项目批量导入回归：模板下载 → 填写 → 预览（逐行结果）→ 确认 → 读回；重复导入不重复创建。

需要 openpyxl（python3 -m pip install openpyxl）。只能对测试库执行：会新建并删除院校、项目。
"""
import io, json, os, sys, uuid, urllib.request, urllib.parse
from openpyxl import load_workbook

B = os.environ.get("ABR_BASE", "http://localhost:8090")
ok = bad = 0


def check(name, cond, info=""):
    global ok, bad
    if cond:
        ok += 1; print("PASS", name)
    else:
        bad += 1; print("FAIL", name, info)


def call(method, path, body=None, token=None, raw=False):
    req = urllib.request.Request(B + path, method=method, data=None if body is None else json.dumps(body).encode(),
                                 headers={"Content-Type": "application/json", **({"Authorization": "Bearer " + token} if token else {})})
    with urllib.request.urlopen(req) as r:
        data = r.read()
        return data if raw else json.loads(data.decode())


def upload(path, content, token):
    boundary = uuid.uuid4().hex
    body = (f"--{boundary}\r\nContent-Disposition: form-data; name=\"file\"; filename=\"import.xlsx\"\r\n"
            f"Content-Type: application/vnd.openxmlformats-officedocument.spreadsheetml.sheet\r\n\r\n").encode() + content + f"\r\n--{boundary}--\r\n".encode()
    req = urllib.request.Request(B + path, method="POST", data=body,
                                 headers={"Content-Type": "multipart/form-data; boundary=" + boundary, "Authorization": "Bearer " + token})
    with urllib.request.urlopen(req) as r:
        return json.loads(r.read().decode())


t = call("POST", "/login", {"username": "admin", "password": "admin123"})["token"]
tpl = call("POST", "/abroad/program/importTemplate", {}, t, raw=True)
wb = load_workbook(io.BytesIO(tpl))
ws = wb.active
headers = [c.value for c in ws[1]]
check("模板 28 列且含关键列", len(headers) == 28 and "项目英文名" in headers and "来源链接" in headers, headers)
check("模板字典列带下拉", len(ws.data_validations.dataValidation) >= 7, len(ws.data_validations.dataValidation))
col = {h: i + 1 for i, h in enumerate(headers)}

SCHOOL = "Test University of Import " + uuid.uuid4().hex[:6]
URL = "https://example.edu/msc"


def fill(rows):
    for r in range(ws.max_row, 1, -1):
        ws.delete_rows(r)
    for i, row in enumerate(rows, start=2):
        for k, v in row.items():
            ws.cell(row=i, column=col[k], value=v)
    buf = io.BytesIO(); wb.save(buf); return buf.getvalue()


base = {"国家": "英国", "院校英文名": SCHOOL, "院校中文名": "导入测试大学", "项目英文名": "MSc Finance", "学位": "硕士", "学科大类": "商科"}
rows = [
    {**base, "专业方向": "金融 / 金融工程", "学制（月）": 12, "入学季": "秋季", "学费（每年）": 38000, "币种": "gbp", "需要APS": "否", "难度档": "D2 很难",
     "申请季": "2027-fall", "均分要求-985": 80, "均分要求-211": 85, "雅思总分": 6.5, "雅思小分": 6.0, "截止日": "滚动", "来源链接": URL},   # 2 create
    {**base, "申请季": "2028-fall", "雅思总分": 7, "来源链接": URL},                                                                  # 3 add_requirement（文件内新项目）
    {**base, "申请季": "2027-fall", "雅思总分": 7, "来源链接": URL},                                                                  # 4 skip 同申请季
    {**base, "项目英文名": "MSc Accounting", "学科大类": "business"},                                                                  # 5 create（复用新院校）
    {**base, "项目英文名": "MSc Marketing", "申请季": "2027-fall", "均分要求-211": 85},                                                # 6 error 缺来源
    {**base, "国家": "法国", "项目英文名": "MSc X"},                                                                                   # 7 error 国家
    {**base, "院校英文名": "No Chinese Name Univ", "院校中文名": None, "项目英文名": "MSc Y"},                                         # 8 error 新院校无中文名
    {**base, "项目英文名": "MSc Accounting", "学科大类": "business"},                                                                  # 9 skip 文件内重复
    {**base, "项目英文名": "MSc Z", "申请季": "2027-fall", "截止日": "2027/01/15", "来源链接": URL},                                    # 10 error 截止日
    {**base, "项目英文名": "MSc W", "申请季": "2027-fall", "雅思总分": 6, "雅思小分": 6.5, "来源链接": URL},                            # 11 error 小分>总分
]
pv = upload("/abroad/program/importPreview", fill(rows), t)
check("预览成功", pv["code"] == 200, pv)
res = {r["rowNo"]: r for r in pv["data"]["rows"]}
expect = {2: "create", 3: "add_requirement", 4: "skip", 5: "create", 6: "error", 7: "error", 8: "error", 9: "skip", 10: "error", 11: "error"}
for no, act in expect.items():
    check(f"第 {no} 行 → {act}", res[no]["action"] == act, res[no])
check("错误行带原因", "来源链接" in "".join(res[6]["messages"]) and "法国" in "".join(res[7]["messages"]), [res[6]["messages"], res[7]["messages"]])
check("预览汇总", pv["data"]["summary"] == {"create": 2, "add_requirement": 1, "skip": 2, "error": 5}, pv["data"]["summary"])
check("预览不写库", call("GET", "/abroad/school/list?schoolName=" + urllib.parse.quote(SCHOOL), None, t)["total"] == 0)

cf = call("POST", f"/abroad/program/importConfirm/{pv['data']['token']}", {}, t)
check("确认导入", cf["code"] == 200 and cf["data"] == {"schoolsCreated": 1, "programsCreated": 2, "requirementsCreated": 2, "skipped": 7}, cf)
again = call("POST", f"/abroad/program/importConfirm/{pv['data']['token']}", {}, t)
check("同一 token 不能重复确认", again["code"] == 500, again)

sl = call("GET", "/abroad/school/list?schoolName=" + urllib.parse.quote(SCHOOL), None, t)
check("新建院校 1 所", sl["total"] == 1 and sl["rows"][0]["schoolName"] == "导入测试大学", sl)
sid = sl["rows"][0]["schoolId"]
pl = call("GET", f"/abroad/program/list?schoolId={sid}", None, t)
check("新建项目 2 个", pl["total"] == 2, pl)
fin = [p for p in pl["rows"] if p["programNameEn"] == "MSc Finance"][0]
check("项目字段读回（字典标签已转值）", fin["degree"] == "master" and fin["fieldGroup"] == "business" and fin["majorDirection"] == "finance"
      and fin["intakes"] == "fall" and fin["currency"] == "GBP" and fin["difficulty"] == "D2" and fin["apsRequired"] == "N" and fin["tuition"] == 38000, fin)
reqs = call("GET", f"/abroad/requirement/list?programId={fin['programId']}", None, t)["data"]
check("申请要求 2 个申请季", sorted(r["season"] for r in reqs) == ["2027-fall", "2028-fall"], reqs)
r27 = [r for r in reqs if r["season"] == "2027-fall"][0]
check("均分、雅思、截止日与来源读回", json.loads(r27["minGpa"]) == {"985": 80, "211": 85} and r27["ieltsMinSub"] == 6.0
      and json.loads(r27["deadlines"])["rolling"] is True and {s["fieldName"] for s in r27["sources"]} == {"minGpa", "ieltsOverall", "ieltsMinSub", "deadlines"}, r27)
check("导入后项目已有核对时间", fin["verifiedTime"] is not None, fin)

pv2 = upload("/abroad/program/importPreview", fill(rows[:5] + [{**base, "申请季": "2029-fall", "雅思总分": 7, "来源链接": URL}]), t)
res2 = {r["rowNo"]: r for r in pv2["data"]["rows"]}
check("重复导入：已存在项目与申请季全部跳过", all(res2[n]["action"] in ("skip", "error") for n in (2, 3, 4, 5, 6)), [res2[n]["action"] for n in res2])
check("重复导入：已存在项目的新申请季 → 补充要求", res2[7]["action"] == "add_requirement" and res2[7]["programId"] == fin["programId"], res2[7])
cf2 = call("POST", f"/abroad/program/importConfirm/{pv2['data']['token']}", {}, t)
check("重复导入不重复创建", cf2["data"]["schoolsCreated"] == 0 and cf2["data"]["programsCreated"] == 0 and cf2["data"]["requirementsCreated"] == 1, cf2)

bad_file = upload("/abroad/program/importPreview", b"not an excel", t)
check("非 Excel 文件给出提示", bad_file["code"] == 500 and "xlsx" in bad_file["msg"], bad_file)

call("DELETE", "/abroad/program/" + ",".join(str(p["programId"]) for p in pl["rows"]), None, t)
call("DELETE", f"/abroad/school/{sid}", None, t)
print(f"\n{ok} passed, {bad} failed")
sys.exit(1 if bad else 0)
